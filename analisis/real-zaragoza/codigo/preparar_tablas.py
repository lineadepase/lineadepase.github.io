"""
Tablas «para humanos»: los CSV del proyecto listos para abrir en Excel y un
diccionario que explica qué es cada columna.

Uso (desde la raíz del repositorio):
    python analisis/real-zaragoza/codigo/preparar_tablas.py

Por qué hace falta esto
-----------------------
Los CSV de datos/entrada están pensados para que los lea el código: separados
por comas, con punto decimal y en UTF-8 sin BOM. Eso es lo estándar, pero si
los abres con doble clic en un Excel configurado en español pasan dos cosas:

    1. Excel no sabe que el archivo está en UTF-8 y lo lee como si fuera
       Windows-1252 (ANSI). Por eso la «ñ» sale como «Ã±» y la «á» como «Ã¡».
       La solución es guardar el archivo con BOM (Byte Order Mark), tres bytes
       al principio (EF BB BF) que le dicen a Excel «esto es UTF-8». En Python
       eso es encoding="utf-8-sig".
    2. En España el separador de listas de Excel es el punto y coma y el
       separador decimal es la coma. Con un CSV separado por comas, Excel mete
       toda la fila en una sola columna. Así que en las copias para Excel uso
       sep=";" y decimal=",".

Los originales NO se tocan: el resto de scripts (y el de R) siguen leyendo los
CSV de datos/entrada tal cual. Todo lo que genera este archivo va a una carpeta
aparte, datos/tablas/, y se puede borrar y volver a crear cuando se quiera.

Qué genera en datos/tablas/
---------------------------
    - datos_real_zaragoza.xlsx  Un Excel con una hoja «Léeme», una hoja
                                «Diccionario» (qué es cada columna) y una hoja
                                por tabla, con cabecera en azul, filtros,
                                columnas a medida y los enlaces clicables.
    - csv_excel/*.csv           Las mismas tablas en CSV listas para Excel en
                                español (UTF-8 con BOM, «;» y coma decimal).
    - DICCIONARIO.md            El diccionario en Markdown. GitHub lo enseña
                                como una tabla bonita sin descargar nada.
"""

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Rutas
# ---------------------------------------------------------------------------
PROYECTO = Path(__file__).resolve().parents[1]
DATOS = PROYECTO / "datos"
DESTINO = DATOS / "tablas"

# Colores del Zaragoza para la cabecera del Excel (en hexadecimal, sin «#»,
# que es como los quiere openpyxl)
AZUL = "213E96"
AZUL_CLARO = "E6EAF7"
BLANCO = "FFFFFF"

# ---------------------------------------------------------------------------
# Diccionario de datos
# ---------------------------------------------------------------------------
# Para cada tabla: la ruta del CSV (relativa a datos/), una frase que explica
# qué contiene y la descripción de cada columna. Lo escribo a mano a propósito:
# es la documentación del proyecto y quiero que esté en un único sitio.
# Si añado una columna a un CSV y me olvido de describirla aquí, el script me
# avisa (ver comprobar_diccionario más abajo).
TABLAS = {
    "partidos": {
        "archivo": "entrada/partidos.csv",
        "que_es": "Un partido oficial por fila (liga y copa) con el resultado final y al descanso.",
        "columnas": {
            "partido_id": "Identificador del partido. En liga coincide con el número de jornada; en copa empieza en 101.",
            "competicion": "Competición: «Primera Federación» o «Copa del Rey».",
            "jornada": "Número de jornada de liga. Vacío en copa.",
            "fecha": "Fecha del partido (año-mes-día).",
            "sede": "«Casa» si el Zaragoza juega de local y «Fuera» si es visitante.",
            "rival": "Nombre del rival.",
            "goles_favor": "Goles del Real Zaragoza.",
            "goles_contra": "Goles del rival.",
            "gf_descanso": "Goles del Zaragoza al descanso.",
            "gc_descanso": "Goles del rival al descanso.",
            "fuente_url": "Página de la que saqué el dato.",
        },
    },
    "goles": {
        "archivo": "entrada/goles.csv",
        "que_es": "Un gol por fila, a favor o en contra, con minuto, goleador y asistencia.",
        "columnas": {
            "partido_id": "Partido al que pertenece el gol (ver tabla partidos).",
            "sentido": "«Favor» si lo marca el Zaragoza y «Contra» si lo marca el rival.",
            "jugador": "Goleador. En los goles en propia puerta, el jugador que la mete.",
            "minuto": "Minuto del gol sin el descuento (un gol en el 45+7 tiene 45 aquí).",
            "anadido": "Minutos de descuento (el 7 del 45+7). 0 si no hay.",
            "parte": "1 = primera parte, 2 = segunda parte.",
            "penalti": "«si» si el gol es de penalti.",
            "propia_puerta": "«si» si es en propia puerta.",
            "asistencia": "Jugador que da el último pase. Vacío si no consta: no significa que no la haya.",
            "fuente_url": "Página de la que saqué el dato.",
        },
    },
    "alineaciones": {
        "archivo": "entrada/alineaciones.csv",
        "que_es": "Once inicial del Zaragoza en cada partido: un titular por fila.",
        "columnas": {
            "partido_id": "Partido (ver tabla partidos).",
            "jugador": "Titular, escrito como en la web del club.",
            "fuente_url": "Página de la que saqué el dato.",
        },
    },
    "cambios": {
        "archivo": "entrada/cambios.csv",
        "que_es": "Sustituciones del Zaragoza: quién entra, quién sale y en qué minuto.",
        "columnas": {
            "partido_id": "Partido (ver tabla partidos).",
            "minuto": "Minuto del cambio.",
            "entra": "Jugador que entra.",
            "sale": "Jugador que sale.",
            "fuente_url": "Página de la que saqué el dato.",
        },
    },
    "tarjetas": {
        "archivo": "entrada/tarjetas.csv",
        "que_es": "Tarjetas que ven los jugadores del Zaragoza.",
        "columnas": {
            "partido_id": "Partido (ver tabla partidos).",
            "jugador": "Jugador amonestado o expulsado.",
            "tarjeta": "«Amarilla», «Roja» o «Segunda amarilla».",
            "minuto": "Minuto base de la tarjeta; vacío si no se ha recogido. Las expulsiones cierran los minutos jugados, sin descuento.",
            "fuente_url": "Página de la que saqué el dato.",
        },
    },
    "tiros": {
        "archivo": "entrada/tiros.csv",
        "que_es": "Tiros y posesión de cada partido, con una nota cuando el desglose no cuadra.",
        "columnas": {
            "partido_id": "Partido (ver tabla partidos).",
            "tiros_zaragoza": "Tiros totales del Zaragoza.",
            "tiros_rival": "Tiros totales del rival.",
            "a_puerta_zaragoza": "Tiros a puerta del Zaragoza.",
            "a_puerta_rival": "Tiros a puerta del rival.",
            "posesion_pct": "Posesión del Zaragoza, en %.",
            "completo": "«si» si a puerta + fuera + bloqueados suman el total; «no» si no cuadra.",
            "fuente": "Fuente que me quedo para ese partido.",
            "nota": "Explicación de la comprobación de los tiros.",
            "fuente_url": "Página de la que saqué el dato.",
        },
    },
    "estadisticas": {
        "archivo": "entrada/estadisticas.csv",
        "que_es": "Otras estadísticas de partido de fuentes públicas, en formato largo (una métrica por fila).",
        "columnas": {
            "partido_id": "Partido (ver tabla partidos).",
            "periodo": "«Partido», «1.ª parte» o «2.ª parte».",
            "metrica": "Nombre de la métrica (posesión, córners, faltas…).",
            "zaragoza": "Valor del Zaragoza.",
            "valor_rival": "Valor del rival.",
            "calidad": "«ok» si lo he podido contrastar; otro valor si hay dudas.",
            "fuente_url": "Página de la que saqué el dato.",
        },
    },
    "sofascore": {
        "archivo": "entrada/sofascore.csv",
        "que_es": "Estadísticas de Sofascore de cada partido, xG incluido, en formato largo (una métrica por fila).",
        "columnas": {
            "partido_id": "Partido (ver tabla partidos).",
            "evento_sofascore": "Identificador del partido en Sofascore.",
            "metrica": "Métrica. Las que acaban en _ok son aciertos y las que acaban en _total, intentos.",
            "zaragoza": "Valor del Zaragoza.",
            "valor_rival": "Valor del rival.",
            "fuente_url": "Dirección de la API de Sofascore de la que copié el dato.",
        },
    },
    "grupo2_partidos": {
        "archivo": "entrada/grupo2_partidos.csv",
        "que_es": "Todos los partidos del grupo 2 con estadísticas de Sofascore. Una fila por equipo y partido: cada partido aparece dos veces, una desde cada lado.",
        "columnas": {
            "jornada": "Jornada de liga.",
            "evento_sofascore": "Identificador del partido en Sofascore.",
            "fecha": "Fecha del partido (AAAA-MM-DD).",
            "equipo": "Equipo al que se refiere la fila (mismo nombre que en la clasificación).",
            "rival": "Rival en ese partido.",
            "sede": "Local o Visitante.",
            "goles_favor": "Goles del equipo.",
            "goles_contra": "Goles del rival.",
            "xg": "Goles esperados (xG) según Sofascore. Vacío si Sofascore no lo da.",
            "xg_rival": "Lo mismo, del rival.",
            "ocasiones_claras": "Ocasiones claras (big chances).",
            "ocasiones_claras_rival": "Lo mismo, del rival.",
            "tiros": "Tiros totales.",
            "tiros_rival": "Lo mismo, del rival.",
            "tiros_a_puerta": "Tiros entre los tres palos.",
            "tiros_a_puerta_rival": "Lo mismo, del rival.",
            "tiros_dentro_area": "Tiros desde dentro del área.",
            "tiros_dentro_area_rival": "Lo mismo, del rival.",
            "posesion_pct": "Posesión, en %.",
            "posesion_pct_rival": "Lo mismo, del rival.",
            "pases": "Pases intentados.",
            "pases_rival": "Lo mismo, del rival.",
            "pases_buenos": "Pases completados.",
            "pases_buenos_rival": "Lo mismo, del rival.",
            "corners": "Saques de esquina.",
            "corners_rival": "Lo mismo, del rival.",
            "faltas": "Faltas cometidas.",
            "faltas_rival": "Lo mismo, del rival.",
            "amarillas": "Tarjetas amarillas.",
            "amarillas_rival": "Lo mismo, del rival.",
            "rojas": "Tarjetas rojas.",
            "rojas_rival": "Lo mismo, del rival.",
            "duelos_ganados_pct": "Duelos ganados, en %.",
            "duelos_ganados_pct_rival": "Lo mismo, del rival.",
            "entradas_ultimo_tercio": "Entradas al último tercio del campo.",
            "entradas_ultimo_tercio_rival": "Lo mismo, del rival.",
            "recuperaciones": "Recuperaciones de balón.",
            "recuperaciones_rival": "Lo mismo, del rival.",
            "paradas": "Paradas del portero.",
            "paradas_rival": "Lo mismo, del rival.",
            "fuente_url": "Dirección de la API de Sofascore de la que copié el dato.",
        },
    },
    "clasificacion": {
        "archivo": "entrada/clasificacion.csv",
        "que_es": "Clasificación del grupo 2 después de cada jornada apuntada.",
        "columnas": {
            "tras_jornada": "Jornada después de la cual se toma la clasificación.",
            "posicion": "Puesto en la tabla.",
            "equipo": "Equipo.",
            "pj": "Partidos jugados.",
            "g": "Partidos ganados.",
            "e": "Partidos empatados.",
            "p": "Partidos perdidos.",
            "gf": "Goles a favor.",
            "gc": "Goles en contra.",
            "puntos": "Puntos.",
            "fuente_url": "Página de la que saqué el dato.",
        },
    },
    "proximos": {
        "archivo": "entrada/proximos.csv",
        "que_es": "Próximos partidos del calendario.",
        "columnas": {
            "jornada": "Jornada.",
            "competicion": "Competición.",
            "fecha": "Fecha prevista (puede cambiar).",
            "sede": "«Casa» o «Fuera».",
            "rival": "Rival.",
            "fuente_url": "Página de la que saqué el dato.",
        },
    },
    "partidos_rivales": {
        "archivo": "entrada/partidos_rivales.csv",
        "que_es": "Resultados del próximo rival esta temporada, para la previa.",
        "columnas": {
            "equipo": "Equipo del que son los resultados.",
            "jornada": "Jornada.",
            "fecha": "Fecha.",
            "local": "Equipo local.",
            "visitante": "Equipo visitante.",
            "goles_local": "Goles del local.",
            "goles_visitante": "Goles del visitante.",
            "fuente_url": "Página de la que saqué el dato.",
        },
    },
    "incidencias": {
        "archivo": "entrada/incidencias.csv",
        "que_es": "Diferencias entre fuentes y qué dato me quedo. Es el registro de decisiones del proyecto.",
        "columnas": {
            "jornada": "Jornada afectada (0 si afecta a todo el proyecto).",
            "campo": "Dato o tabla afectada.",
            "problema": "Qué no cuadra.",
            "decision": "Qué he hecho con ello.",
        },
    },
    "goles_video": {
        "archivo": "entrada/video/goles_video.csv",
        "que_es": "Cada gol de liga apuntado viendo los resúmenes en YouTube: cómo nace, por dónde llega y desde dónde se remata.",
        "columnas": {
            "gol_id": "Jornada y número de gol del partido (6-02 = segundo gol de la J6).",
            "partido_id": "Partido (ver tabla partidos).",
            "jornada": "Jornada de liga.",
            "rival": "Rival.",
            "sede": "«Casa» o «Fuera».",
            "sentido": "«Favor» o «Contra».",
            "minuto": "Minuto en texto, con el descuento (45+7).",
            "marcador": "Resultado justo después del gol, siempre Zaragoza-rival.",
            "goleador": "Quién marca.",
            "asistencia": "Quién da el último pase, si consta.",
            "inicio": "Cómo empieza la jugada: ataque posicional, transición, balón parado, penalti o sin dato.",
            "detalle_inicio": "La misma idea en una frase corta.",
            "carril": "Izquierda, centro o derecha DEL EQUIPO QUE ATACA.",
            "ultimo_pase": "Tipo de último pase: centro lateral, pase atrás, pase raso por dentro, al espacio, rechace, córner, falta…",
            "zona_remate": "Zona del último toque: área pequeña, área grande (centro o lateral), fuera del área o penalti.",
            "remate": "Con qué se remata: pie, cabeza o propia puerta.",
            "revision_fvs": "Si se ve una revisión en el monitor (FVS).",
            "defensores_area": "Defensores dentro del área en el momento del último pase, solo si se ven todos.",
            "atacantes_area": "Atacantes dentro del área en ese mismo momento.",
            "tiro_x": "Posición estimada del remate, en metros hacia la portería rival (0 a 105).",
            "tiro_y": "Posición estimada del remate, en metros desde la banda derecha del que ataca (0 a 68).",
            "pase_x": "Posición estimada del último pase (x).",
            "pase_y": "Posición estimada del último pase (y).",
            "segundos_video": "Dónde está la jugada en el resumen, en segundos.",
            "confianza": "Alta, media o baja: lo bien que se ve la jugada.",
            "nota": "Lo que se ve en el vídeo, explicado en una o dos frases.",
        },
    },
    "resumenes_youtube": {
        "archivo": "entrada/video/resumenes_youtube.csv",
        "que_es": "Vídeos de resumen que he usado para apuntar los goles.",
        "columnas": {
            "partido_id": "Partido (ver tabla partidos).",
            "jornada": "Jornada.",
            "rival": "Rival.",
            "canal": "Canal de YouTube que publica el resumen.",
            "titulo": "Título del vídeo.",
            "duracion": "Duración (minutos:segundos).",
            "url": "Enlace al vídeo.",
        },
    },
    "jugadores": {
        "archivo": "salida/jugadores.csv",
        "que_es": "Tabla calculada: minutos y participación en goles de cada jugador en liga.",
        "columnas": {
            "jugador": "Jugador.",
            "partidos": "Partidos jugados.",
            "titularidades": "Partidos como titular.",
            "minutos": "Minutos jugados, sin contar el descuento.",
            "goles": "Goles.",
            "asistencias": "Asistencias (solo las que constan).",
            "participaciones_con_minutos": "Goles + asistencias.",
            "participaciones_por_90": "Goles + asistencias por cada 90 minutos jugados.",
            "minutos_por_participacion": "Minutos que tarda de media en participar en un gol. Vacío si no ha participado en ninguno.",
        },
    },
    "minutos_partido": {
        "archivo": "salida/minutos_partido.csv",
        "que_es": "Tabla calculada: minutos de cada jugador en cada partido de liga.",
        "columnas": {
            "partido_id": "Partido (ver tabla partidos).",
            "jugador": "Jugador.",
            "rol": "«Titular» o «Suplente».",
            "entra": "Minuto en el que entra (0 si es titular).",
            "sale": "Minuto en el que sale (90 si acaba el partido).",
            "minutos": "Minutos jugados.",
        },
    },
}


# ---------------------------------------------------------------------------
# Lectura y comprobaciones
# ---------------------------------------------------------------------------

def leer(archivo):
    """Lee un CSV del proyecto. Uso utf-8-sig porque lee bien tanto los archivos
    con BOM (los de datos/salida) como los que no lo llevan (los de entrada)."""
    return pd.read_csv(DATOS / archivo, encoding="utf-8-sig")


def enteros_limpios(df):
    """pandas lee como decimales (float) las columnas de enteros que tienen algún
    hueco vacío, porque el vacío (NaN) es un float. Resultado: un 7 se guarda como
    «7,0». Aquí paso esas columnas al tipo «Int64» de pandas, que admite enteros
    con huecos, y así se escriben limpias."""
    df = df.copy()
    for columna in df.columns:
        valores = df[columna].dropna()
        if pd.api.types.is_float_dtype(df[columna]) and len(valores) and (valores % 1 == 0).all():
            df[columna] = df[columna].astype("Int64")
    return df


def comprobar_diccionario(tablas):
    """Comprueba que cada columna de cada CSV tiene su descripción y que no
    describo columnas que ya no existen. Si algo falla, paro el script: prefiero
    enterarme ahora que tener un diccionario que miente."""
    problemas = []
    for nombre, info in TABLAS.items():
        reales = list(tablas[nombre].columns)
        descritas = list(info["columnas"])
        sin_describir = [c for c in reales if c not in descritas]
        sobrantes = [c for c in descritas if c not in reales]
        if sin_describir:
            problemas.append(f"{nombre}: faltan por describir {sin_describir}")
        if sobrantes:
            problemas.append(f"{nombre}: el diccionario habla de columnas que no existen {sobrantes}")
    if problemas:
        raise SystemExit("El diccionario no cuadra con los CSV:\n  " + "\n  ".join(problemas))


def tipo_legible(serie):
    """Traduce el tipo de pandas (dtype) a algo que entienda cualquiera."""
    if pd.api.types.is_integer_dtype(serie):
        return "Número entero"
    if pd.api.types.is_float_dtype(serie):
        # Una columna de enteros con huecos vacíos la lee pandas como float;
        # si todos los valores que hay son enteros, la sigo llamando entero.
        valores = serie.dropna()
        return "Número entero" if len(valores) and (valores % 1 == 0).all() else "Número decimal"
    return "Texto"


def ejemplo(serie):
    """Primer valor no vacío de la columna, recortado para que no ocupe media hoja."""
    valores = serie.dropna()
    if valores.empty:
        return ""
    texto = str(valores.iloc[0])
    return texto if len(texto) <= 60 else texto[:57] + "…"


def tabla_diccionario(tablas):
    """Junta el diccionario en una sola tabla: tabla, columna, descripción, tipo y ejemplo."""
    filas = []
    for nombre, info in TABLAS.items():
        datos = tablas[nombre]
        for columna, descripcion in info["columnas"].items():
            filas.append({"Tabla": nombre, "Columna": columna, "Qué es": descripcion,
                          "Tipo": tipo_legible(datos[columna]), "Ejemplo": ejemplo(datos[columna])})
    return pd.DataFrame(filas)


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

def dar_formato(hoja, df, anchos_max=60):
    """Formato común a todas las hojas: cabecera azul con letra blanca, fila de
    cabecera fija al hacer scroll, filtros y ancho de columna según el contenido."""
    cabecera = Font(bold=True, color=BLANCO)
    relleno = PatternFill("solid", fgColor=AZUL)
    franja = PatternFill("solid", fgColor=AZUL_CLARO)
    linea = Border(bottom=Side(style="thin", color="D6E0F0"))

    for celda in hoja[1]:
        celda.font = cabecera
        celda.fill = relleno
        celda.alignment = Alignment(vertical="center", wrap_text=True)
    hoja.row_dimensions[1].height = 30

    # Filas alternas en azul muy claro, que ayudan a no perderse en tablas largas
    for i, fila in enumerate(hoja.iter_rows(min_row=2), start=2):
        for celda in fila:
            celda.border = linea
            celda.alignment = Alignment(vertical="top", wrap_text=True)
            if i % 2 == 1:
                celda.fill = franja

    # Ancho de cada columna: el texto más largo entre la cabecera y los valores,
    # con un mínimo para que no quede apretado y un máximo para las notas largas
    for j, columna in enumerate(df.columns, start=1):
        largo = max([len(str(columna))] + [len(str(v)) for v in df[columna].dropna()])
        hoja.column_dimensions[get_column_letter(j)].width = min(max(10, largo + 2), anchos_max)

    hoja.freeze_panes = "A2"                 # la cabecera se queda fija
    hoja.auto_filter.ref = hoja.dimensions   # flechas de filtro en la cabecera


def enlaces_clicables(hoja, df):
    """Convierte las columnas de URL en hipervínculos de verdad (se pueden pulsar)."""
    for j, columna in enumerate(df.columns, start=1):
        if "url" not in columna:
            continue
        for i in range(2, len(df) + 2):
            celda = hoja.cell(row=i, column=j)
            if isinstance(celda.value, str) and celda.value.startswith("http"):
                celda.hyperlink = celda.value
                celda.font = Font(color="1F4FBF", underline="single")


def hoja_desde_tabla(libro, titulo, df, anchos_max=60):
    """Añade una hoja al libro con los datos de un DataFrame y le da formato.
    Los valores vacíos (NaN) los escribo como celdas vacías, no como «nan»."""
    hoja = libro.create_sheet(title=titulo[:31])   # Excel no admite nombres de hoja de más de 31 caracteres
    hoja.append(list(df.columns))
    for fila in df.itertuples(index=False):
        hoja.append([None if pd.isna(v) else v for v in fila])
    dar_formato(hoja, df, anchos_max)
    enlaces_clicables(hoja, df)
    return hoja


def crear_excel(tablas, diccionario, ruta):
    libro = Workbook()

    # Hoja 1: «Léeme», una portada que explica qué hay en el archivo
    leeme = libro.active
    leeme.title = "Léeme"
    leeme["A1"] = "Real Zaragoza 2026/27 · Datos del proyecto"
    leeme["A1"].font = Font(bold=True, size=16, color=AZUL)
    leeme["A2"] = "Línea de pase · Ángel Marqueta Alonso. Datos públicos; cada fila lleva su fuente."
    leeme["A4"] = "Hoja"
    leeme["B4"] = "Qué contiene"
    leeme["C4"] = "Filas"
    for celda in (leeme["A4"], leeme["B4"], leeme["C4"]):
        celda.font = Font(bold=True, color=BLANCO)
        celda.fill = PatternFill("solid", fgColor=AZUL)
    fila = 5
    leeme.cell(row=fila, column=1, value="Diccionario")
    leeme.cell(row=fila, column=2, value="Qué significa cada columna de cada tabla, con su tipo y un ejemplo.")
    for nombre, info in TABLAS.items():
        fila += 1
        leeme.cell(row=fila, column=1, value=nombre)
        leeme.cell(row=fila, column=2, value=info["que_es"])
        leeme.cell(row=fila, column=3, value=len(tablas[nombre]))
    leeme.column_dimensions["A"].width = 22
    leeme.column_dimensions["B"].width = 100
    leeme.column_dimensions["C"].width = 8

    # Hoja 2: el diccionario
    hoja_desde_tabla(libro, "Diccionario", diccionario, anchos_max=90)

    # Una hoja por tabla
    for nombre in TABLAS:
        hoja_desde_tabla(libro, nombre, tablas[nombre])

    libro.save(ruta)


# ---------------------------------------------------------------------------
# Markdown para GitHub
# ---------------------------------------------------------------------------

def crear_markdown(tablas, diccionario, ruta):
    """Escribe el diccionario como tablas Markdown, una por tabla de datos.
    El carácter «|» separa columnas en Markdown, así que lo escapo si aparece en un texto."""
    lineas = ["# Diccionario de datos · Real Zaragoza 2026/27", "",
              "Qué es cada tabla del proyecto y qué significa cada columna. "
              "Este archivo lo genera `codigo/preparar_tablas.py`; no lo edites a mano.", ""]
    for nombre, info in TABLAS.items():
        lineas += [f"## {nombre}", "", f"{info['que_es']} Archivo: `datos/{info['archivo']}` "
                   f"({len(tablas[nombre])} filas).", "",
                   "| Columna | Qué es | Tipo | Ejemplo |", "|---|---|---|---|"]
        # to_dict("records") me da cada fila como diccionario; así puedo usar la
        # columna «Qué es» por su nombre aunque tenga tilde y espacio
        for f in diccionario[diccionario["Tabla"] == nombre].to_dict("records"):
            ej = str(f["Ejemplo"]).replace("|", "\\|")
            lineas.append(f"| `{f['Columna']}` | {f['Qué es']} | {f['Tipo']} | {ej} |")
        lineas.append("")
    ruta.write_text("\n".join(lineas), encoding="utf-8")


# ---------------------------------------------------------------------------
# Programa principal
# ---------------------------------------------------------------------------

def main():
    # 1. Leo todas las tablas que describe el diccionario
    tablas = {nombre: enteros_limpios(leer(info["archivo"])) for nombre, info in TABLAS.items()}
    comprobar_diccionario(tablas)
    diccionario = tabla_diccionario(tablas)

    # 2. Copias en CSV para Excel en español: UTF-8 con BOM, punto y coma y coma decimal
    carpeta_csv = DESTINO / "csv_excel"
    carpeta_csv.mkdir(parents=True, exist_ok=True)
    for nombre, df in tablas.items():
        df.to_csv(carpeta_csv / f"{nombre}.csv", index=False, sep=";", decimal=",", encoding="utf-8-sig")
    diccionario.to_csv(carpeta_csv / "diccionario.csv", index=False, sep=";", encoding="utf-8-sig")

    # 3. El Excel con formato y el diccionario en Markdown
    crear_excel(tablas, diccionario, DESTINO / "datos_real_zaragoza.xlsx")
    crear_markdown(tablas, diccionario, DESTINO / "DICCIONARIO.md")

    print(f"{len(tablas)} tablas y {len(diccionario)} columnas descritas. Todo en {DESTINO}")


if __name__ == "__main__":
    main()
